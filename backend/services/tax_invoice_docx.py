"""Tax Invoice as a Word document (11 Sep 2026, user request).

Same sections, same order and same figures as the on-screen sheet / PDF:
header (seller block + TAX INVOICE meta), buyer & shipping cards, the
service table (billing-unit columns when the invoice came from a
timesheet), GST summary + totals, amount in words, bank details +
declaration with seal and QR, and the website-only footer.

Built with python-docx so Finance can open it in Word and tweak wording
without touching the figures. A4 portrait, navy/white table headers, no
external assets beyond the bundled logo/seal PNGs.
"""
from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services import tax_invoice as ti

NAVY = "173B7A"
NAVY_RGB = RGBColor(0x17, 0x3B, 0x7A)
MUTED_RGB = RGBColor(0x64, 0x74, 0x8B)
BORDER = "D9E2EC"


# ---------------------------------------------------------------- helpers

def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _borders(table, color: str = BORDER, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell_text(cell, text: str, *, bold: bool = False, size: float = 8.5,
               color: RGBColor | None = None, align=None, italic: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text or "")
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _add_line(cell, text: str, *, bold: bool = False, size: float = 8.5,
              color: RGBColor | None = None, align=None, italic: bool = False):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text or "")
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def _hyperlink(paragraph, url: str, text: str, color: str = "FFFFFF") -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    b = OxmlElement("w:b")
    r_pr.append(b)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "17")
    r_pr.append(sz)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    link.append(new_run)
    paragraph._p.append(link)


def _set_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)


def _money(v: Any) -> str:
    return ti.format_inr(ti.num(v))


def _plain(v: Any) -> str:
    return ti.format_inr(ti.num(v)).replace("INR ", "")


def _qty(v: Any) -> str:
    f = ti.num(v)
    return f"{int(f):,}" if f == int(f) else f"{f:,.2f}".rstrip("0").rstrip(".")


def _gap(doc, pts: float = 3) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)
    r = p.add_run(" ")
    r.font.size = Pt(2)


# ---------------------------------------------------------------- builder

def build_invoice_docx(inv: ti.Invoice) -> bytes:
    totals = ti.compute_totals(inv)
    seller = inv.seller or ti.seller_from_settings()
    bank = inv.bank or ti.bank_from_settings()
    cols = inv.columns

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.1)
    sec.top_margin = Cm(0.8)
    sec.bottom_margin = Cm(0.8)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(8.5)
    base.paragraph_format.space_after = Pt(0)
    base.paragraph_format.space_before = Pt(0)

    usable = 21.0 - 2.2  # cm

    # ---- header -----------------------------------------------------------
    hdr = doc.add_table(rows=1, cols=2)
    hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(hdr)
    left, right = hdr.rows[0].cells
    left.text = ""
    p0 = left.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    if ti.LOGO_PATH.exists():
        try:
            p0.add_run().add_picture(str(ti.LOGO_PATH), height=Cm(1.1))
        except Exception:
            p0.add_run("KARNEX").bold = True
    else:
        p0.add_run("KARNEX").bold = True
    _add_line(left, (seller.get("name") or "").upper(), bold=True, size=10, color=NAVY_RGB)
    if seller.get("tagline"):
        _add_line(left, seller["tagline"], italic=True, size=8, color=NAVY_RGB)
    addr = ", ".join(x for x in (seller.get("address_line1"), seller.get("address_line2")) if x)
    _add_line(left, addr, size=8, color=MUTED_RGB)
    _add_line(left, f"State Name: {seller.get('state') or '—'}    State Code: {seller.get('state_code') or '—'}",
              size=8, color=MUTED_RGB)
    _add_line(left, f"Email: {seller.get('email') or seller.get('contact_email') or '—'}", size=8, color=MUTED_RGB)
    _add_line(left, f"CIN No.: {seller.get('cin') or '—'}", size=8, color=MUTED_RGB)

    right.text = ""
    pr = right.paragraphs[0]
    pr.paragraph_format.space_after = Pt(4)
    run = pr.add_run("TAX INVOICE")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY_RGB
    for label, value in (("Invoice No.", inv.invoice_no), ("Invoice Date", inv.invoice_date),
                         ("P.O. No.", inv.po_no or "—"), ("P.O. Date", inv.po_date or "—"),
                         ("GSTIN No.", seller.get("gstin") or "—"), ("PAN No.", seller.get("pan") or "—")):
        p = _add_line(right, "", size=8.5)
        r1 = p.add_run(f"{label:<13}: ")
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = MUTED_RGB
        r2 = p.add_run(str(value or "—"))
        r2.font.size = Pt(8.5)
        r2.bold = True
    _set_widths(hdr, [usable * 0.55, usable * 0.45])

    _gap(doc)

    # ---- buyer / shipping cards ------------------------------------------
    cards = doc.add_table(rows=2, cols=2)
    _borders(cards)
    for idx, title in enumerate(("Buyer Details", "Shipping Details")):
        c = cards.rows[0].cells[idx]
        _shade(c, NAVY)
        _cell_text(c, title, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    b = inv.buyer
    bc = cards.rows[1].cells[0]
    _cell_text(bc, b.name or "—", bold=True, size=9)
    _add_line(bc, b.address or "—", size=8)
    _add_line(bc, f"GSTIN: {b.gstn or '—'} · PAN: {b.pan or '—'}", size=8)
    _add_line(bc, f"State: {b.state_name or '—'} ({b.state_code or '—'})", size=8)
    sc = cards.rows[1].cells[1]
    _cell_text(sc, b.shipping or b.name or "—", bold=True, size=9)
    _add_line(sc, b.shipping_address or b.address or "—", size=8)
    _add_line(sc, f"GSTIN: {b.gstn or '—'} · PAN: {b.pan or '—'}", size=8)
    _add_line(sc, f"State: {b.state_name or '—'} ({b.state_code or '—'})", size=8)
    _set_widths(cards, [usable / 2, usable / 2])

    _gap(doc)

    # ---- service table ----------------------------------------------------
    if cols:
        heads = ["S.No", "Description of Services", "SAC Code", cols["cost"], cols["qty"],
                 cols["leave"], cols["per_day"], cols["amount"]]
        widths = [1.0, 5.8, 1.5, 2.2, 1.5, 1.4, 2.6, 2.8]
    else:
        heads = ["S.No", "Description of Services", "SAC Code", inv.qty_label, inv.rate_label, "Amount"]
        widths = [1.2, 7.6, 2.0, 2.4, 2.6, 2.8]
    items = inv.items or []
    n_rows = max(len(items), 3) + 1
    svc = doc.add_table(rows=n_rows, cols=len(heads))
    _borders(svc)
    for i, h in enumerate(heads):
        c = svc.rows[0].cells[i]
        _shade(c, NAVY)
        _cell_text(c, h, bold=True, size=7, color=RGBColor(0xFF, 0xFF, 0xFF),
                   align=WD_ALIGN_PARAGRAPH.RIGHT if i >= 3 else None)
    for r in range(1, n_rows):
        it = items[r - 1] if r - 1 < len(items) else None
        cells = svc.rows[r].cells
        if it is None:
            for c in cells:
                _cell_text(c, "", size=8)
            continue
        desc = ti.line_description(it.employee_name, getattr(it, "service_month", None) or None)
        _cell_text(cells[0], str(r), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(cells[1], desc, size=8)
        if getattr(it, "period_label", ""):
            _add_line(cells[1], f"Billing period {it.period_label}", size=7.5, color=MUTED_RGB)
        _cell_text(cells[2], it.sac or ti.default_sac(), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        amt = ti.line_amount(it)
        if cols:
            cost = it.monthly_cost if it.monthly_cost is not None else it.rate_per_hour
            _cell_text(cells[3], _plain(cost), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[4], _qty(it.billing_hours), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[5], _qty(it.leave_days), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[6], _plain(it.rate_per_day) if it.rate_per_day is not None else "—",
                       size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[7], _plain(amt), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True)
        else:
            _cell_text(cells[3], _qty(it.billing_hours), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[4], _money(it.rate_per_hour), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _cell_text(cells[5], _money(amt), size=8, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True)
    _set_widths(svc, widths)

    _gap(doc)

    # ---- GST summary + totals ---------------------------------------------
    gst = doc.add_table(rows=1, cols=2)
    _no_borders(gst)
    lcell, rcell = gst.rows[0].cells
    _cell_text(lcell, "GST Summary", bold=True, size=9, color=NAVY_RGB)
    left_rows = [("CGST (9%)", totals.cgst), ("SGST (9%)", totals.sgst), ("IGST (18%)", totals.igst),
                 ("Total GST", totals.total_gst)]
    lt = lcell.add_table(rows=len(left_rows), cols=2)
    _borders(lt)
    for i, (k, v) in enumerate(left_rows):
        _cell_text(lt.rows[i].cells[0], k, size=8.5, bold=(i == len(left_rows) - 1))
        _cell_text(lt.rows[i].cells[1], _money(v), size=8.5, align=WD_ALIGN_PARAGRAPH.RIGHT,
                   bold=(i == len(left_rows) - 1))
    _set_widths(lt, [5.0, 3.6])

    rcell.text = ""
    right_rows = [("Sub Total", totals.subtotal), ("CGST (9%)", totals.cgst), ("SGST (9%)", totals.sgst),
                  ("IGST (18%)", totals.igst), ("Total GST", totals.total_gst), ("Grand Total", totals.total)]
    rt = rcell.add_table(rows=len(right_rows), cols=2)
    _borders(rt)
    for i, (k, v) in enumerate(right_rows):
        last = i == len(right_rows) - 1
        c0, c1 = rt.rows[i].cells
        if last:
            _shade(c0, NAVY)
            _shade(c1, NAVY)
        white = RGBColor(0xFF, 0xFF, 0xFF) if last else None
        _cell_text(c0, k, size=9 if last else 8.5, bold=last, color=white)
        _cell_text(c1, _money(v), size=9 if last else 8.5, bold=last, color=white, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_widths(rt, [5.0, 3.6])
    _set_widths(gst, [usable / 2, usable / 2])

    _gap(doc)

    # ---- amount in words ---------------------------------------------------
    words = doc.add_table(rows=1, cols=1)
    _borders(words)
    wc = words.rows[0].cells[0]
    _shade(wc, "F8FAFC")
    _cell_text(wc, "AMOUNT CHARGEABLE (IN WORDS)", bold=True, size=7.5, color=MUTED_RGB)
    _add_line(wc, totals.amount_in_words, bold=True, italic=True, size=9)
    _add_line(wc, "TAX AMOUNT (IN WORDS)", bold=True, size=7.5, color=MUTED_RGB)
    _add_line(wc, totals.tax_in_words, bold=True, italic=True, size=9)

    _gap(doc)

    # ---- bank + declaration -----------------------------------------------
    bottom = doc.add_table(rows=1, cols=2)
    _borders(bottom)
    bk, dc = bottom.rows[0].cells
    _cell_text(bk, "Bank Details", bold=True, size=9, color=NAVY_RGB)
    for k, v in (("Bank", bank.get("bank_name")), ("Account Name", bank.get("account_name")),
                 ("A/C No.", bank.get("account_number")), ("IFSC", bank.get("ifsc")),
                 ("Branch", bank.get("branch")), ("Account Type", bank.get("account_type")),
                 ("SWIFT", bank.get("swift_code")), ("UPI", bank.get("upi_id"))):
        if not v and k in ("SWIFT", "UPI"):
            continue
        p = _add_line(bk, "", size=8)
        r1 = p.add_run(f"{k}: ")
        r1.font.size = Pt(7.5)
        r1.font.color.rgb = MUTED_RGB
        r2 = p.add_run(str(v or "—"))
        r2.font.size = Pt(8.5)
        r2.bold = True
    _cell_text(dc, "Declaration", bold=True, size=9, color=NAVY_RGB)
    _add_line(dc, inv.footer_text or ti.DEFAULT_FOOTER, size=8)
    _add_line(dc, "", size=4)
    _add_line(dc, seller.get("signatory_line") or "For Karnex Software Solutions Pvt. Ltd.", bold=True, size=8.5)
    pic = _add_line(dc, "", size=8)
    if ti.SEAL_PATH.exists():
        try:
            pic.add_run().add_picture(str(ti.SEAL_PATH), height=Cm(1.6))
        except Exception:
            pass
    qr_embedded = False
    if getattr(inv, "share_url", None):
        try:
            from services.invoice_share import qr_png_bytes
            png = qr_png_bytes(inv.share_url)
            if png:
                pic.add_run("    ")
                pic.add_run().add_picture(io.BytesIO(png), height=Cm(1.6))
                qr_embedded = True
        except Exception:
            qr_embedded = False
    sig = _add_line(dc, "Authorized Signatory", size=8, color=MUTED_RGB)
    if getattr(inv, "share_url", None) and not qr_embedded:
        r = sig.add_run("   ·   ")
        r.font.size = Pt(7.5)
        r.font.color.rgb = MUTED_RGB
        _hyperlink(sig, inv.share_url, "view this invoice online", color="173B7A")
    _set_widths(bottom, [usable / 2, usable / 2])

    _gap(doc)

    # ---- footer: website only, clickable -----------------------------------
    foot = doc.add_table(rows=1, cols=1)
    _borders(foot, color=NAVY)
    fc = foot.rows[0].cells[0]
    _shade(fc, NAVY)
    fc.text = ""
    fp = fc.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(3)
    fp.paragraph_format.space_after = Pt(3)
    url = (seller.get("footer_website_url") or "").strip()
    if url and not url.lower().startswith(("http://", "https://")):
        url = "https://" + url
    label = (seller.get("website") or url.replace("https://", "").replace("http://", "")).strip() or "www.karnex.in"
    if url:
        _hyperlink(fp, url, label)
    else:
        r = fp.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    extra = (seller.get("footer_text") or "").strip()
    if extra:
        r = fp.add_run(f"   ·   {extra}")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_filename(inv: ti.Invoice) -> str:
    return ti.pdf_filename(inv).replace(".pdf", ".docx")
