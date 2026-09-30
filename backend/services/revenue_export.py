"""Period-close workbook for the CEO revenue report (18 Sep 2026; v3 21 Sep 2026).

One .xlsx, one sheet per section of `services.revenue_report.revenue_report`,
so the board pack is the same numbers the page shows — nothing recomputed
here. Written with openpyxl's normal workbook (small: a few hundred cells).
The workbook follows the page's zoom and filters — a quarter view exports the
quarter, a customer-filtered view exports that customer only.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '₹#,##0;[Red]-₹#,##0'
PCT = '0.0"%"'


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], formats: dict[int, str] | None = None):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for idx, fmt in (formats or {}).items():
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=idx).number_format = fmt
    for col in range(1, len(headers) + 1):
        width = max(len(str(headers[col - 1])), *(len(str(r[col - 1])) for r in rows if len(r) >= col)) if rows else len(headers[col - 1])
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(12, width + 2))
    ws.freeze_panes = "A2"
    return ws


def _n(v):
    return v if v is not None else ""


def build_revenue_workbook(report: dict, placements: dict | None = None) -> bytes:
    h, t, m, c, f, a = (report["headline"], report["targets"], report["margin"], report["collections"],
                        report["forecast"], report["ageing"])
    wb = Workbook()
    wb.remove(wb.active)

    summary = [
        ["Period", report["label"]],
        ["Zoom", {"month": "Month", "quarter": "Quarter", "fy": "Financial year"}.get(
            report.get("period", "month"), "Month")],
        ["From", report.get("period_start", "")], ["To", report.get("period_end", "")],
        ["As of", report["as_of"]],
        ["Revenue billed (excl. GST)", h["billed"]], ["GST", h["gst"]], ["Billed incl. GST", h["billed_incl_gst"]],
        ["Invoices raised", h["invoices"]], ["Cash collected", h["collected"]],
        ["Collection rate %", _n(h["collection_rate_pct"])],
        [f"Billed previous period ({h.get('comparison_label', 'vs last month')})", h["billed_prev"]],
        ["Change %", _n(h["billed_mom_pct"])],
        ["Billed same period last year", h["billed_yoy"]], ["YoY %", _n(h["billed_yoy_pct"])],
        ["Outstanding receivables", h["outstanding"]], ["Overdue", a["overdue_total"]],
        ["Overdue 90+ days", a["overdue_90_plus"]],
        ["Period target", _n(t["month_target"])], ["Attainment %", _n(t["month_attainment_pct"])],
        ["Run-rate", t["run_rate"]],
        [f"{t['fy_label']} target", _n(t["fy_target"])], [f"{t['fy_label']} billed to date", t["fy_billed_to_date"]],
        [f"{t['fy_label']} projection", t["fy_projection"]],
        ["People cost (CTC/12)", m["cost"]], ["Gross margin", m["gross_margin"]],
        ["Gross margin %", _n(m["gross_margin_pct"])], ["Heads without CTC", m["heads_without_ctc"]],
        ["DSO (days)", _n(c["dso_days"])], ["Avg days to pay", _n(c["avg_days_to_pay"])],
        *([["Expected cash, next 3 months", report["cashflow"]["expected_total_at_pace"]],
           ["People cost, next 3 months", report["cashflow"]["people_cost_total"]],
           ["Net cash position", report["cashflow"]["net_total_at_pace"]],
           ["Avg days paid late", report["cashflow"]["slip_days"]]]
          if report.get("cashflow") else []),
        ["Deployed headcount", report["efficiency"]["deployed_heads"]],
        ["Revenue per head", _n(report["efficiency"]["revenue_per_head"])],
        ["Approved, not invoiced", report["pipeline"]["approved_uninvoiced"]["amount"]],
        ["Active PO balance", report["pipeline"]["active_po_balance"]],
        ["PO cover (months)", _n(report["pipeline"]["po_cover_months"])],
        [f"Forecast next {len(f['months'])} months", f["total"]],
    ]
    ws = _sheet(wb, "Summary", ["Metric", "Value"], summary)
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=2)
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            label = str(ws.cell(row=r, column=1).value)
            cell.number_format = PCT if "%" in label else "0.0" if "days" in label.lower() or "months" in label.lower() else MONEY

    _sheet(wb, "Alerts", ["Level", "Alert", "Detail"],
           [[x["level"], x["title"], x["detail"]] for x in report["alerts"]])
    trend_title = {"month": "12-month trend", "quarter": "8-quarter trend",
                   "fy": "5-year trend"}.get(report.get("period", "month"), "Trend")
    _sheet(wb, trend_title, ["Period", "Billed", "Collected", "Invoices"],
           [[s["label"], s["billed"], s["collected"], s["invoices"]] for s in report["series"]],
           {2: MONEY, 3: MONEY})
    _sheet(wb, "By customer",
           ["Customer", "Billed", "Share %", "Invoices", "Collected", "Previous", "Change %"],
           [[r["customer"], r["billed"], r["share_pct"], r["invoices"], r["collected"],
             r.get("previous", ""), _n(r.get("change_pct"))]
            for r in report["by_customer"]["rows"]], {2: MONEY, 3: PCT, 5: MONEY, 6: MONEY, 7: PCT})
    bp = report.get("by_project") or {"rows": []}
    _sheet(wb, "By project",
           ["Project", "Customer", "Billed", "Share %", "Invoices", "People cost", "Margin",
            "Margin %", "Heads"],
           [[r["project"], r["customer"], r["billed"], r["share_pct"], r["invoices"],
             r["cost"], r["margin"], _n(r["margin_pct"]), r["heads"]] for r in bp["rows"]],
           {3: MONEY, 4: PCT, 6: MONEY, 7: MONEY, 8: PCT})
    be = report.get("by_employee") or {"rows": []}
    _sheet(wb, "By employee",
           ["Employee", "Emp ID", "Customer", "Project", "Billed", "Invoices", "People cost",
            "Margin", "Margin %", "Deployed"],
           [[r["employee"], r.get("employee_code") or "", r.get("customer") or "",
             r.get("project") or "", r["billed"], r["invoices"], r["cost"], r["margin"],
             _n(r["margin_pct"]), "Yes" if r.get("deployed") else "No"] for r in be["rows"]]
           + ([["Not linked to a timesheet", "", "", "", be.get("unlinked_billed", 0),
                be.get("unlinked_invoices", 0), "", "", "", ""]]
              if be.get("unlinked_billed") else []),
           {5: MONEY, 7: MONEY, 8: MONEY, 9: PCT})
    _sheet(wb, "Margin", ["Customer", "Billed", "People cost", "Margin", "Margin %", "Heads"],
           [[r["customer"], r["billed"], r["cost"], r["margin"], _n(r["margin_pct"]), r["heads"]]
            for r in m["by_customer"]], {2: MONEY, 3: MONEY, 4: MONEY, 5: PCT})
    _sheet(wb, "Ageing", ["Bucket", "Invoices", "Amount"],
           [[b["label"], b["count"], b["amount"]] for b in a["buckets"]]
           + [["", "", ""], ["Top overdue", "", ""]]
           + [[x["customer"], "", x["overdue"]] for x in a["top_overdue_customers"]], {3: MONEY})
    _sheet(wb, "Forecast", ["Month", "Amount", "Heads", "Roll-offs", "Within PO cover"],
           [[x["label"], x["amount"], x["heads"], x["roll_off_count"], "Yes" if x["po_covered"] else "No"]
            for x in f["months"]] + [["Assumptions", f["assumptions"], "", "", ""]], {2: MONEY})
    dims = report["dimensions"]
    rows = []
    for title, key in (("Opportunity type", "by_type"), ("Branch", "by_branch"), ("Sales owner", "by_owner")):
        for r in dims[key]:
            rows.append([title, r["label"], r["billed"], r["share_pct"], r["invoices"]])
    _sheet(wb, "Dimensions", ["Dimension", "Value", "Billed", "Share %", "Invoices"], rows, {3: MONEY, 4: PCT})
    cf = report.get("cashflow")
    if cf:
        _sheet(wb, "Cash flow",
               ["Month", "Expected (on terms)", "Expected (at our pace)", "Invoices",
                "People cost", "Net (on terms)", "Net (at our pace)"],
               [[m["label"], m["expected"], m["expected_at_pace"], m["invoices"],
                 m["people_cost"], m["net"], m["net_at_pace"]] for m in cf["months"]]
               + [["", "", "", "", "", "", ""],
                  ["Total", cf["expected_total"], cf["expected_total_at_pace"], "",
                   cf["people_cost_total"], cf["net_total"], cf["net_total_at_pace"]],
                  ["Overdue now", cf["overdue"]["amount"], "", cf["overdue"]["count"], "", "", ""],
                  ["Approved, not invoiced", cf["unbilled_ready"]["amount"], "",
                   cf["unbilled_ready"]["count"], "", "", ""],
                  ["Due beyond the horizon", cf["beyond_horizon"]["amount"], "",
                   cf["beyond_horizon"]["count"], "", "", ""],
                  ["Avg days late", cf["slip_days"], "", "", "", "", ""],
                  ["Assumptions", cf["assumptions"], "", "", "", "", ""]],
               {2: MONEY, 3: MONEY, 5: MONEY, 6: MONEY, 7: MONEY})
        _sheet(wb, "Expected inflows",
               ["Customer", "Invoice", "Amount", "Due", "Expected (at our pace)", "Days overdue"],
               [[r["customer"], r["number"], r["amount"], r["due"], r["expected"],
                 r["overdue_days"] or ""] for r in cf["top_expected"]], {3: MONEY})
    lk = report["leakage"]
    _sheet(wb, "Leakage", ["Metric", "Value"],
           [["Approved timesheets", lk["approved_sheets"]], ["LOP days", lk["lop_days"]],
            ["Covered by weekend work", lk["lop_covered_by_weekend_work_days"]],
            ["No-billing period days", lk["no_billing_period_days"]]])

    if placements:
        _placement_sheets(wb, placements)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _placement_sheets(wb: Workbook, pl: dict) -> None:
    """Internal vs external placements (25 Sep 2026) — same window and filters."""
    h, rev = pl["headline"], pl["revenue"]
    _sheet(wb, "Placements", ["Customer", "Internal", "External", "Unknown", "Total", "Internal %",
                              "Billed · internal", "Billed · external", "Billed · unknown"],
           [[r["customer"], r["internal"], r["external"], r["unknown"], r["total"],
             _n(r["internal_pct"]), r["billed_internal"], r["billed_external"], r["billed_unknown"]]
            for r in pl["by_customer"]]
           + [["All customers", h["internal"], h["external"], h["unknown"], h["placements"],
               _n(h["internal_pct"]), rev["internal"], rev["external"], rev["unknown"]]],
           {6: PCT, 7: MONEY, 8: MONEY, 9: MONEY})
    _sheet(wb, "Placement list", ["Placed on", "Employee", "Emp ID", "Customer", "Project", "Type",
                                  "Karnex joined", "Why", "Billed in period"],
           [[r["placed_on"], r["employee"], _n(r["employee_code"]), _n(r["customer"]), r["project"],
             r["kind"].title(), _n(r["karnex_joined"]), r["reason"], r["billed"]] for r in pl["rows"]],
           {9: MONEY})
