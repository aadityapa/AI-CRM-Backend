"""Serve CRM-uploaded files (resumes, CVs, certificates, invoices, timesheet attachments)."""
from __future__ import annotations

import mimetypes
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse

from crm_deps import CurrentUser, any_crm_role
from services.crm_common import INLINE_SAFE_EXTENSIONS, resolve_crm_file

router = APIRouter(prefix="/api/crm-files", tags=["CRM: Files"])

# Ensure common office types are registered (stdlib misses some on Windows).
mimetypes.add_type("application/pdf", ".pdf")
mimetypes.add_type(
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx",
)
mimetypes.add_type("application/msword", ".doc")


@router.get("/{rel_path:path}")
def get_crm_file(rel_path: str, preview: str | None = None,
                 user: CurrentUser = Depends(any_crm_role)):
    path = resolve_crm_file(rel_path)
    if path is None:
        raise HTTPException(status_code=404, detail="File not found")
    ext = path.suffix.lower()

    # In-app spreadsheet preview (27 Aug 2026): ?preview=xlsx renders the
    # workbook as escaped HTML tables server-side (openpyxl is already here),
    # so imported timesheets open right in the preview modal — no SheetJS on
    # the client, no download round-trip. Values only, escaped, JSON-wrapped;
    # the frontend runs it through its sanitizer before rendering.
    if preview == "xlsx" and ext in (".xlsx", ".xlsm"):
        from html import escape as _esc

        from fastapi.responses import JSONResponse
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            raise HTTPException(status_code=422, detail="Not a readable Excel file")
        MAX_SHEETS, MAX_ROWS, MAX_COLS = 5, 300, 50
        parts: list[str] = []
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows_html: list[str] = []
            for r_i, row in enumerate(ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS,
                                                   values_only=True)):
                cells = "".join(
                    f"<td>{_esc('' if v is None else (v.strftime('%d-%m-%Y %H:%M').replace(' 00:00', '') if hasattr(v, 'strftime') else str(v)))}</td>"
                    for v in row)
                rows_html.append(f"<tr>{cells}</tr>")
            if not rows_html:
                continue
            truncated = (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS
            parts.append(
                f"<h3>{_esc(ws.title)}</h3><table>{''.join(rows_html)}</table>"
                + ("<p><em>Preview truncated — download for the full sheet.</em></p>"
                   if truncated else ""))
        wb.close()
        if not parts:
            parts.append("<p>(Empty workbook)</p>")
        return JSONResponse({"success": True, "data": {"html": "".join(parts)}})
    media_type, _ = mimetypes.guess_type(path.name)
    # Only render types that cannot execute in the browser. Anything else — an
    # uploaded .html or .svg, say — downloads instead, so a file supplied by one
    # user can never run as script on this origin in another user's session.
    inline = ext in INLINE_SAFE_EXTENSIONS
    return FileResponse(
        path,
        media_type=media_type or "application/octet-stream",
        filename=Path(path.name).name,
        content_disposition_type="inline" if inline else "attachment",
        headers={"X-Content-Type-Options": "nosniff"},
    )
