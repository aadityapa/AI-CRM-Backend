"""Full data backup (14 Sep 2026): one ZIP with every selected dataset as
Excel / CSV / JSON plus attached files; Admin/CEO only; secrets never leave."""
from __future__ import annotations

import importlib
import io
import json
import sys
import zipfile
from pathlib import Path

import pytest
import sqlalchemy as sa
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.dialects.postgresql import ARRAY, INET, JSONB, UUID
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


@compiles(JSONB, "sqlite")
def _j(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(ARRAY, "sqlite")
def _a(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(UUID, "sqlite")
def _u(e, c, **k):  # noqa: ANN001
    return "VARCHAR(36)"


@compiles(INET, "sqlite")
def _i(e, c, **k):  # noqa: ANN001
    return "VARCHAR(64)"


for _m in ["base", "rbac", "customers", "opportunities", "projects", "leave", "timesheets",
           "finance", "hr", "candidates", "masters", "requirements", "profiles", "resumes",
           "ai_links", "scheduling", "user_profiles", "template_requests"]:
    importlib.import_module(f"models.{_m}")

import crm_deps  # noqa: E402
from models import Base, Candidate, Customer  # noqa: E402
from models.base import users_table_stub  # noqa: E402
from services import data_backup as backup  # noqa: E402
import routers.crm.backup as backup_router  # noqa: E402


@pytest.fixture()
def db(tmp_path, monkeypatch):
    engine = sa.create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = Session(bind=engine, future=True)
    s.execute(users_table_stub.insert().values(id=1))
    s.add(Customer(name="Acme Motors"))
    # a CV on disk, referenced the way the app does (/api/crm-files/<rel>)
    uploads = tmp_path / "uploads"
    (uploads / "resumes").mkdir(parents=True)
    (uploads / "resumes" / "abc.pdf").write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(backup, "resolve_crm_file",
                        lambda rel: (uploads / rel) if (uploads / rel).is_file() else None)
    s.add(Candidate(first_name="Asha", last_name="Rao", email="asha@example.com",
                    cv_url="/api/crm-files/resumes/abc.pdf"))
    s.add(Candidate(first_name="No", last_name="CV", email="nocv@example.com",
                    cv_url="/api/crm-files/resumes/missing.pdf"))
    s.commit()
    monkeypatch.setattr(backup, "BACKUP_DIR", tmp_path / "backups")
    yield s
    s.close()


def test_every_dataset_table_exists_in_the_orm_or_is_legacy():
    orm = set(Base.metadata.tables)
    for d in backup.DATASETS:
        missing = [t for t in d.tables if t not in orm]
        assert not missing, f"{d.key}: unknown ORM tables {missing}"
    assert len({d.key for d in backup.DATASETS}) == len(backup.DATASETS)
    sheets = [t[:31] for d in backup.DATASETS for t in (*d.tables, *d.legacy_tables)]
    assert len(sheets) == len(set(sheets)), "Excel sheet names (31 chars) must stay unique"


def test_no_orm_table_is_forgotten_by_the_registry():
    """A new table must be assigned to a dataset (or explicitly excluded here)
    — otherwise 'complete backup' silently stops being complete."""
    covered = {t for d in backup.DATASETS for t in (*d.tables, *d.legacy_tables)}
    excluded = set()   # nothing today; list tables here only with a reason
    forgotten = sorted(set(Base.metadata.tables) - covered - excluded)
    assert not forgotten, f"Add these tables to a dataset in services/data_backup.py: {forgotten}"


def test_archive_has_csv_json_excel_files_and_readme(db, tmp_path):
    job = backup.BackupJob(id="t1", datasets=["customers", "candidates"], requested_by="Karan")
    path = backup.build_archive(db, backup.resolve_datasets(["customers", "candidates"]), job,
                                out_dir=tmp_path / "backups")
    assert path.exists() and job.tables_done == job.tables_total
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        assert "README.txt" in names and "karnex-backup.xlsx" in names
        assert "customers/csv/customers.csv" in names and "candidates/json/candidates.json" in names
        rows = json.loads(zf.read("candidates/json/candidates.json"))
        asha = next(r for r in rows if r["email"] == "asha@example.com")
        assert asha["_files_folder"] and asha["_files_folder"].startswith("files/candidates/")
        assert any(n.startswith(asha["_files_folder"] + "/") and n.endswith("abc.pdf") for n in names)
        nocv = next(r for r in rows if r["email"] == "nocv@example.com")
        assert nocv.get("_files_folder") is None
        readme = zf.read("README.txt").decode()
        assert "Missing files: 1" in readme
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(zf.read("karnex-backup.xlsx")), read_only=True)
        assert "candidates" in wb.sheetnames and "customers" in wb.sheetnames
    assert job.files == 1 and job.files_missing == 1
    assert (path.with_suffix(".json")).exists()


def _legacy_tables(db):
    """The interview platform's raw tables, as the app creates them (no ORM)."""
    db.execute(sa.text("DROP TABLE IF EXISTS registration_data"))   # the ORM stub
    db.execute(sa.text("""CREATE TABLE registration_data (
        id INTEGER PRIMARY KEY, full_name TEXT, email TEXT, username TEXT, role TEXT,
        password_hash TEXT, password_salt TEXT, is_active INTEGER)"""))
    db.execute(sa.text("INSERT INTO registration_data VALUES (1,'Karan','k@karnex.in','karan','hr','HASH','SALT',1)"))
    db.execute(sa.text("""CREATE TABLE interview_schedule (
        id TEXT PRIMARY KEY, candidate_email TEXT, access_key TEXT, invite_token TEXT,
        active_device_id TEXT, "select" TEXT)"""))
    db.execute(sa.text("INSERT INTO interview_schedule VALUES ('s1','c@x.com','KEY123','TOK456','dev','kw')"))
    db.execute(sa.text("CREATE TABLE login_data (id INTEGER PRIMARY KEY, secret TEXT)"))
    db.commit()


def test_secrets_are_never_dumped_from_legacy_tables(db):
    _legacy_tables(db)
    cols, rows = backup.dump_table(db, "registration_data")
    assert rows and rows[0]["full_name"] == "Karan" and rows[0]["email"] == "k@karnex.in"
    assert not {"password_hash", "password_salt"} & set(cols)
    cols, rows = backup.dump_table(db, "interview_schedule")
    assert rows[0]["candidate_email"] == "c@x.com" and rows[0]["select"] == "kw"   # reserved word quoted
    assert not {"access_key", "invite_token", "active_device_id"} & set(cols)     # substring redaction
    assert backup.dump_table(db, "login_data") == ([], [])            # denied outright
    assert backup.dump_table(db, "no_such_table") == ([], [])         # skipped, not raised
    assert backup.dump_table(db, 'x"; DROP TABLE y') == ([], [])      # name never reaches SQL


def test_app_settings_credential_values_are_redacted(db):
    from models import AppSetting
    db.add(AppSetting(key="smtp.password", value="hunter2"))
    db.add(AppSetting(key="invoice.sac_code", value="998513"))
    db.commit()
    _, rows = backup.dump_table(db, "app_settings")
    by = {r["key"]: r["value"] for r in rows}
    assert by["smtp.password"] == "<redacted>" and by["invoice.sac_code"] == "998513"


def test_csv_neutralises_formula_injection_but_excel_keeps_text():
    assert backup._csv_cell("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert backup._csv_cell("+919876543210") == "'+919876543210"
    assert backup._csv_cell(-5) == -5
    assert backup._excel_value("+919876543210") == "+919876543210"
    assert backup._excel_value({"a": 1}) == '{"a": 1}'


def test_prune_keeps_only_the_newest_and_drops_stale_partials(tmp_path):
    import os, time
    for i in range(5):
        p = tmp_path / f"karnex-backup-2026-01-0{i + 1}T00-00-00-abc{i}.zip"
        p.write_bytes(b"PK"); p.with_suffix(".json").write_text("{}")
        os.utime(p, (time.time() - (10 - i) * 60, time.time() - (10 - i) * 60))
    stale = tmp_path / "old.zip.part"; stale.write_bytes(b"x")
    os.utime(stale, (time.time() - 8 * 3600,) * 2)
    fresh = tmp_path / "live.zip.part"; fresh.write_bytes(b"x")
    backup._prune(keep=2, out_dir=tmp_path)
    left = sorted(p.name for p in tmp_path.glob("karnex-backup-*.zip"))
    assert len(left) == 2 and left[-1].endswith("abc4.zip")
    assert not stale.exists() and fresh.exists()


def test_failed_build_leaves_no_partial_archive(db, tmp_path, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("disk on fire")
    monkeypatch.setattr(backup, "dump_table", boom)
    job = backup.BackupJob(id="t2", datasets=["customers"], requested_by="x")
    with pytest.raises(RuntimeError):
        backup.build_archive(db, backup.resolve_datasets(["customers"]), job, out_dir=tmp_path / "b")
    assert not list((tmp_path / "b").glob("*.part")) and not list((tmp_path / "b").glob("*.tmp"))


def test_only_one_build_runs_at_a_time(monkeypatch):
    import threading
    gate = threading.Event()

    def slow(db, datasets, job, out_dir=None):
        gate.wait(5)
        return Path("/dev/null")
    monkeypatch.setattr(backup, "build_archive", slow)
    monkeypatch.setattr(backup, "_JOBS", {}); monkeypatch.setattr(backup, "_RUNNING", None)
    ds = backup.resolve_datasets(["settings"])
    first = backup.start_job(ds, "a", lambda: type("S", (), {"close": lambda self: None})())
    with pytest.raises(RuntimeError):
        backup.start_job(ds, "b", lambda: None)
    assert backup.job_status()["id"] == first.id and backup.job_status()["status"] in ("queued", "running")
    gate.set()


def test_resolve_datasets_all_and_unknown():
    assert len(backup.resolve_datasets([])) == len(backup.DATASETS)
    assert len(backup.resolve_datasets(["all"])) == len(backup.DATASETS)
    assert [d.key for d in backup.resolve_datasets(["employees", "finance"])] == ["employees", "finance"]
    with pytest.raises(ValueError):
        backup.resolve_datasets(["nope"])


def test_archive_path_refuses_anything_outside_backup_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(backup, "BACKUP_DIR", tmp_path)
    assert backup.archive_path("../etc/passwd") is None
    assert backup.archive_path("karnex-backup-2026-01-01T00-00-00.zip") is None   # not on disk
    (tmp_path / "karnex-backup-2026-01-01T00-00-00.zip").write_bytes(b"PK")
    assert backup.archive_path("karnex-backup-2026-01-01T00-00-00.zip") is not None


def _client(user):
    app = FastAPI()
    app.include_router(backup_router.router)
    app.dependency_overrides[crm_deps.get_current_user] = lambda: user
    return TestClient(app)


def test_endpoints_are_admin_or_ceo_only():
    sales = crm_deps.CurrentUser(id=2, username="s", full_name="Sales", roles={"Sales"})
    hr = crm_deps.CurrentUser(id=3, username="h", full_name="HR", roles={"HR"})
    for u in (sales, hr):
        c = _client(u)
        assert c.get("/api/admin/backup/datasets").status_code == 403
        assert c.post("/api/admin/backup", json={"datasets": ["all"]}).status_code == 403
        assert c.get("/api/admin/backup/history").status_code == 403
    ceo = crm_deps.CurrentUser(id=1, username="k", full_name="Karan", roles={"CEO"})
    r = _client(ceo).get("/api/admin/backup/datasets")
    assert r.status_code == 200 and r.json()["success"]
    assert {d["key"] for d in r.json()["data"]["datasets"]} >= {"customers", "employees", "ai_hiring"}


def test_start_rejects_unknown_dataset_and_bad_download_name():
    admin = crm_deps.CurrentUser(id=1, username="a", full_name="Admin", roles={"Admin"})
    c = _client(admin)
    assert c.post("/api/admin/backup", json={"datasets": ["bogus"]}).status_code == 400
    assert c.get("/api/admin/backup/download/..%2F..%2Fetc%2Fpasswd").status_code == 404
    assert c.get("/api/admin/backup/download/karnex-backup-nope.zip").status_code == 404


def test_date_window_filters_dated_tables_and_keeps_masters_whole(db, tmp_path):
    from datetime import date, datetime, timezone
    from models import Candidate, Department
    old = db.query(Candidate).filter_by(email="asha@example.com").one()
    old.created_at = datetime(2024, 3, 1, tzinfo=timezone.utc)
    new = db.query(Candidate).filter_by(email="nocv@example.com").one()
    new.created_at = datetime(2026, 5, 20, tzinfo=timezone.utc)
    db.add(Department(name="Engineering"))
    db.commit()
    assert backup.date_column_for("candidates") == "created_at"
    assert backup.date_column_for("departments") is None
    _, rows = backup.dump_table(db, "candidates", date_from=date(2026, 1, 1), date_to=date(2026, 12, 31))
    assert [r["email"] for r in rows] == ["nocv@example.com"]
    _, rows = backup.dump_table(db, "candidates", date_from=date(2024, 1, 1), date_to=date(2024, 12, 31))
    assert [r["email"] for r in rows] == ["asha@example.com"]
    _, deps = backup.dump_table(db, "departments", date_from=date(2030, 1, 1), date_to=date(2030, 12, 31))
    assert len(deps) == 1                                   # masters always whole
    job = backup.BackupJob(id="t3", datasets=["candidates"], requested_by="x",
                           date_from=date(2026, 1, 1), date_to=date(2026, 12, 31))
    path = backup.build_archive(db, backup.resolve_datasets(["candidates"]), job, out_dir=tmp_path / "w")
    with zipfile.ZipFile(path) as zf:
        assert "Period    : 2026-01-01 to 2026-12-31" in zf.read("README.txt").decode()
        assert len(json.loads(zf.read("candidates/json/candidates.json"))) == 1


def test_start_rejects_inverted_window(monkeypatch):
    monkeypatch.setattr(backup, "_JOBS", {}); monkeypatch.setattr(backup, "_RUNNING", None)
    admin = crm_deps.CurrentUser(id=1, username="a", full_name="Admin", roles={"Admin"})
    r = _client(admin).post("/api/admin/backup", json={"datasets": ["all"], "date_from": "2026-05-01", "date_to": "2026-04-01"})
    assert r.status_code == 400
