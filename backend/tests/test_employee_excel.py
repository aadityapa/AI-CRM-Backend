"""Employee Excel import/export (services/employee_excel.py, 26 Aug 2026).

Pins the round trip: template → fill → import creates employees with names
resolved to master ids; bad rows fail alone; existing emails are skipped;
export uses the import's own column layout (plus ID) so it can be re-imported.

Run:  cd backend && python -m pytest tests/test_employee_excel.py -q
"""
from __future__ import annotations

import importlib
import io

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.dialects.postgresql import ARRAY, INET, JSONB, UUID
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool


@compiles(JSONB, "sqlite")
def _j(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(ARRAY, "sqlite")
def _a(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(UUID, "sqlite")
def _u(e, c, **k):  # noqa: ANN001
    return "VARCHAR(36)"


@compiles(INET, "sqlite")
def _i(e, c, **k):  # noqa: ANN001
    return "VARCHAR(64)"


for _m in [
    "base", "rbac", "customers", "opportunities", "projects", "leave", "timesheets",
    "finance", "hr", "candidates", "masters", "requirements", "profiles", "resumes",
    "ai_links", "scheduling", "user_profiles", "template_requests", "access_templates",
]:
    importlib.import_module(f"models.{_m}")

from models.base import Base  # noqa: E402
from models import Department, Designation, Employee  # noqa: E402
from services.employee_excel import (  # noqa: E402
    _HEADERS, build_export_workbook, build_template_workbook, import_workbook,
)


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = Session(bind=engine, future=True)
    from models.base import users_table_stub
    s.execute(users_table_stub.insert().values(id=1))
    s.commit()
    try:
        yield s
    finally:
        s.close()


def _seed_masters(db):
    db.add(Department(name="Engineering"))
    db.add(Designation(name="Software Engineer"))
    db.flush()


def _filled_workbook(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(_HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in _HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_has_all_columns_and_reference_sheet(db):
    _seed_masters(db)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_template_workbook(db)))
    ws = wb["Employees"]
    assert [c.value for c in ws[1]] == _HEADERS
    ref = wb["Reference"]
    assert ref["A1"].value == "Departments"
    assert ref["A2"].value == "Engineering"


def test_import_resolves_names_and_creates(db):
    _seed_masters(db)
    data = _filled_workbook([{
        "First Name": "Asha", "Last Name": "Verma", "Email": "asha@karnex.in",
        "Department": "engineering",       # case-insensitive
        "Designation": "Software Engineer",
        "Profile Type": "internal",
        "Employment Type": "Full Time",    # space normalised to Full_Time
        "Skills": "Java, SQL",
        "Date of Joining": "01/09/2026",
        "Active": "Yes",
    }])
    result = import_workbook(db, data)
    db.commit()
    assert result["failed"] == [] and len(result["created"]) == 1
    emp = db.execute(select(Employee)).scalars().one()
    assert emp.department_id is not None and emp.designation_id is not None
    assert emp.skills == ["Java", "SQL"]
    assert emp.employment_type == "Full_Time"
    assert str(emp.date_of_joining) == "2026-09-01"


def test_bad_row_fails_alone_and_existing_email_skips(db):
    _seed_masters(db)
    db.add(Employee(first_name="Old", email="old@karnex.in"))
    db.commit()
    data = _filled_workbook([
        {"First Name": "Good", "Email": "good@karnex.in"},
        {"First Name": "Bad", "Email": "bad@karnex.in", "Department": "No Such Dept"},
        {"First Name": "Dup", "Email": "old@karnex.in"},
    ])
    result = import_workbook(db, data)
    db.commit()
    assert len(result["created"]) == 1
    assert len(result["failed"]) == 1 and "No Such Dept" in result["failed"][0]["error"]
    assert len(result["skipped"]) == 1 and "already exists" in result["skipped"][0]["reason"]
    # the good row really landed despite its bad neighbour
    assert db.execute(select(Employee).where(Employee.email == "good@karnex.in")).scalars().first()


def test_export_round_trips_into_import(db):
    _seed_masters(db)
    data = _filled_workbook([{"First Name": "Asha", "Email": "asha@karnex.in",
                              "Department": "Engineering"}])
    import_workbook(db, data)
    db.commit()
    exported = build_export_workbook(db)
    # The export (with its extra ID column) must be an acceptable import file:
    # same people → all skipped as existing, nothing failed.
    result = import_workbook(db, exported)
    assert result["failed"] == []
    assert len(result["skipped"]) == 1
