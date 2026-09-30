"""Internal vs external placements (25 Sep 2026) — `services.placements_report`.

Every customer placement is labelled from facts already recorded: an earlier
placement anywhere, or a Karnex joining date more than 30 days before the
customer onboarding → internal; joined within 30 days → external (new hire);
no joining date and no earlier placement → unknown. In-memory SQLite.
"""
from __future__ import annotations

import importlib
from datetime import date

import pytest
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import ARRAY, INET, JSONB, UUID
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool


@compiles(JSONB, "sqlite")
def _j(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(ARRAY, "sqlite")
def _a(e, c, **k):  # noqa: ANN001
    return "JSON"


@compiles(UUID, "sqlite")
def _u(e, c, **k):  # noqa: ANN001
    return "CHAR(36)"


@compiles(INET, "sqlite")
def _i(e, c, **k):  # noqa: ANN001
    return "VARCHAR(45)"


for _m in [
    "base", "rbac", "customers", "opportunities", "projects", "leave", "timesheets",
    "finance", "hr", "candidates", "masters", "requirements", "profiles", "resumes",
    "ai_links", "scheduling", "user_profiles", "template_requests", "access_templates",
]:
    importlib.import_module(f"models.{_m}")

from models.base import Base  # noqa: E402
from models import (  # noqa: E402
    Customer, Employee, Invoice, PaymentStatus, Project, ProjectEmployee, Timesheet, TimesheetStatus,
)
from services import placements_report as pr  # noqa: E402

TODAY = date(2026, 9, 25)


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = Session(bind=engine, future=True)
    from models.base import users_table_stub
    s.execute(users_table_stub.insert().values(id=1))
    s.commit()
    try:
        yield s
    finally:
        s.close()


def _place(db, emp, proj, on):
    pe = ProjectEmployee(project_id=proj.id, employee_id=emp.id, billing_rate=100000, onboarding_date=on)
    db.add(pe); db.flush()
    return pe


def _bill(db, emp, proj, on, amount, kind="Tax", sheet_month=None):
    ts = Timesheet(project_id=proj.id, employee_id=emp.id, month=sheet_month or on.month, year=on.year,
                   status=TimesheetStatus.APPROVED)
    db.add(ts); db.flush()
    db.add(Invoice(invoice_number=f"INV-{ts.id}-{kind}", project_id=proj.id, invoice_date=on,
                   sub_total=amount, tax_amount=0, grand_total=amount, paid_amount=0,
                   balance_amount=amount, payment_status=PaymentStatus.UNPAID,
                   timesheet_id=ts.id, kind=kind))
    db.flush()


@pytest.fixture()
def book(db):
    harman = Customer(name="HARMAN"); bosch = Customer(name="BOSCH")
    db.add_all([harman, bosch]); db.flush()
    ph = Project(customer_id=harman.id, name="Infotainment")
    pb = Project(customer_id=bosch.id, name="Braking")
    db.add_all([ph, pb]); db.flush()

    def emp(name, doj):
        e = Employee(first_name=name, email=f"{name.lower()}@karnex.in", date_of_joining=doj)
        db.add(e); db.flush()
        return e

    veteran = emp("Veteran", date(2023, 3, 1))        # on rolls for years → internal
    hire = emp("Hire", date(2026, 9, 1))              # joined 4 days before → external
    late = emp("Late", date(2026, 9, 12))             # joined AFTER onboarding → external
    mover = emp("Mover", date(2026, 8, 20))           # new hire in Aug at BOSCH, moved in Sep → internal
    ghost = emp("Ghost", None)                        # no joining date → unknown
    edge = emp("Edge", date(2026, 8, 6))              # exactly 30 days before → external
    _place(db, veteran, ph, date(2026, 9, 10))
    _place(db, hire, ph, date(2026, 9, 5))
    _place(db, late, pb, date(2026, 9, 8))
    _place(db, mover, pb, date(2026, 8, 25))
    _place(db, mover, ph, date(2026, 9, 15))
    _place(db, ghost, pb, date(2026, 9, 20))
    _place(db, edge, ph, date(2026, 9, 5))
    # Revenue in September, through timesheets — plus a Proforma that must not count
    # and a manual invoice nobody can be attributed to.
    _bill(db, veteran, ph, date(2026, 9, 30), 300000)
    _bill(db, hire, ph, date(2026, 9, 30), 100000)
    _bill(db, hire, ph, date(2026, 9, 30), 999999, kind="Proforma", sheet_month=8)
    db.add(Invoice(invoice_number="MANUAL-1", project_id=pb.id, invoice_date=date(2026, 9, 29),
                   sub_total=50000, tax_amount=0, grand_total=50000, paid_amount=0,
                   balance_amount=50000, payment_status=PaymentStatus.UNPAID))
    db.commit()
    return db


def test_classify_rules():
    joined = date(2026, 1, 1)
    assert pr.classify(date(2026, 1, 31), joined, None)[0] == pr.EXTERNAL      # 30 days = still a hire
    assert pr.classify(date(2026, 2, 1), joined, None)[0] == pr.INTERNAL       # 31 days = on rolls
    assert pr.classify(date(2025, 12, 20), joined, None)[0] == pr.EXTERNAL     # joined after onboarding
    assert pr.classify(date(2026, 2, 1), None, None)[0] == pr.UNKNOWN
    prior = pr.Placement(1, 1, "X", None, 1, "BOSCH", 9, "Braking", date(2025, 12, 1), None)
    kind, reason, _ = pr.classify(date(2026, 1, 5), joined, prior)
    assert kind == pr.INTERNAL and "Braking" in reason and "BOSCH" in reason


def test_september_split_and_every_row_explains_itself(book):
    r = pr.placements_report(book, month_raw="2026-09", today=TODAY)
    h = r["headline"]
    assert h["placements"] == 6 and h["unique_people"] == 6
    assert (h["internal"], h["external"], h["unknown"]) == (2, 3, 1)
    assert h["internal_pct"] == 40.0 and h["external_pct"] == 60.0
    kinds = {row["employee"]: row["kind"] for row in r["rows"]}
    assert kinds == {"Veteran": "internal", "Mover": "internal", "Hire": "external",
                     "Late": "external", "Edge": "external", "Ghost": "unknown"}
    assert all(row["reason"] for row in r["rows"])
    # Newest first.
    assert r["rows"][0]["employee"] == "Ghost"


def test_a_previous_placement_at_another_customer_makes_a_move_internal_even_when_filtered(book):
    harman = book.query(Customer).filter_by(name="HARMAN").one()
    r = pr.placements_report(book, month_raw="2026-09", customer_id=harman.id, today=TODAY)
    mover = next(row for row in r["rows"] if row["employee"] == "Mover")
    assert mover["kind"] == "internal" and "BOSCH" in mover["reason"]
    assert r["headline"]["placements"] == 4    # Veteran, Hire, Mover, Edge at HARMAN


def test_revenue_split_excludes_proformas_and_reports_manual_invoices(book):
    rev = pr.placements_report(book, month_raw="2026-09", today=TODAY)["revenue"]
    assert rev["internal"] == 300000 and rev["external"] == 100000
    assert rev["unattributed"] == 50000 and rev["total"] == 450000
    assert rev["internal_pct"] == 75.0


def test_by_customer_counts_and_billing(book):
    rows = {r["customer"]: r for r in pr.placements_report(book, month_raw="2026-09", today=TODAY)["by_customer"]}
    assert (rows["HARMAN"]["internal"], rows["HARMAN"]["external"]) == (2, 2)
    assert rows["HARMAN"]["billed_internal"] == 300000 and rows["HARMAN"]["billed_external"] == 100000
    assert (rows["BOSCH"]["external"], rows["BOSCH"]["unknown"]) == (1, 1)


def test_quarter_series_and_previous_period(book):
    r = pr.placements_report(book, month_raw="2026-09", period_raw="quarter", today=TODAY)
    assert r["window"]["key"] == "2026-Q2" and len(r["series"]) == 8
    last = r["series"][-1]
    assert last["label"] == "Q2 26" and last["internal"] + last["external"] + last["unknown"] == 7
    # The August BOSCH placement is the Mover's first → a new hire.
    assert r["headline"]["external"] == 4


def test_month_comparison_counts_the_previous_month(book):
    h = pr.placements_report(book, month_raw="2026-09", today=TODAY)["headline"]
    assert h["previous"]["placements"] == 1 and h["previous"]["external"] == 1


def test_custom_range_buckets_by_week_then_by_month(book):
    r = pr.placements_report(book, date_from="2026-09-01", date_to="2026-09-14", today=TODAY)
    assert r["window"]["kind"] == "custom" and r["window"]["bucket"] == "week"
    assert [b["key"] for b in r["series"]] == ["2026-09-01", "2026-09-07", "2026-09-14"]
    assert r["headline"]["placements"] == 4    # 5, 5, 8, 10 Sep
    long = pr.placements_report(book, date_from="2026-04-01", date_to="2026-09-30", today=TODAY)
    assert long["window"]["bucket"] == "month" and len(long["series"]) == 6


def test_bad_ranges_are_refused():
    for args in ({"date_from": "2026-09-10"}, {"date_from": "2026-09-10", "date_to": "2026-09-01"},
                 {"date_from": "2020-01-01", "date_to": "2026-01-01"}, {"date_from": "10/09/2026", "date_to": "2026-09-30"}):
        with pytest.raises(ValueError):
            pr.resolve_window(None, None, args.get("date_from"), args.get("date_to"), TODAY)


def test_undated_assignments_are_counted_not_guessed(book):
    e = Employee(first_name="Nodate", email="nodate@karnex.in", date_of_joining=date(2026, 1, 1))
    book.add(e); book.flush()
    proj = book.query(Project).first()
    book.add(ProjectEmployee(project_id=proj.id, employee_id=e.id, billing_rate=1))
    book.commit()
    r = pr.placements_report(book, month_raw="2026-09", today=TODAY)
    assert r["headline"]["undated"] == 1 and r["headline"]["placements"] == 6


def test_route_is_admin_only_and_the_workbook_carries_the_split(book):
    import inspect
    from openpyxl import load_workbook
    from io import BytesIO
    from routers.crm import reports
    assert "role_required()" in inspect.getsource(reports.revenue_placements)
    from services.revenue_export import _placement_sheets
    from openpyxl import Workbook
    wb = Workbook()
    _placement_sheets(wb, pr.placements_report(book, month_raw="2026-09", today=TODAY))
    buf = BytesIO(); wb.save(buf)
    names = load_workbook(BytesIO(buf.getvalue())).sheetnames
    assert "Placements" in names and "Placement list" in names
